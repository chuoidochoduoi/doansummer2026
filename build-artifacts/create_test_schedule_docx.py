from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"D:\gitlap\doAnSummer2026\build-artifacts\CareS_Test_Milestones_Table.docx"

rows = [
    ("Create Test Plan", "3", "11/05/2026", "13/05/2026"),
    ("Prepare Test Environment and Data", "4", "14/05/2026", "18/05/2026"),
    ("Design Unit Test Cases", "15", "19/05/2026", "31/07/2026"),
    ("Execute and Complete Unit Test Baseline", "12", "01/06/2026", "14/08/2026"),
    ("Design Integration Test Cases", "6", "24/07/2026", "21/08/2026"),
    ("Execute and Review Integration Test Round 1", "4", "22/08/2026", "23/08/2026"),
    ("Analyze and Correct Integration Issues", "9", "24/08/2026", "09/09/2026"),
    ("Execute Integration Test Round 2", "3", "10/09/2026", "11/09/2026"),
    ("Design System Test Cases", "5", "17/08/2026", "27/08/2026"),
    ("Execute and Review System Test Round 1", "3", "28/08/2026", "29/08/2026"),
    ("Correct System Workflow Issues", "7", "30/08/2026", "09/09/2026"),
    ("Execute System Test Round 2", "2", "10/09/2026", "11/09/2026"),
    ("Execute Final Unit Regression", "3", "12/09/2026", "14/09/2026"),
    ("Prepare Acceptance Test", "3", "12/09/2026", "14/09/2026"),
    ("Execute Acceptance Test", "4", "15/09/2026", "15/09/2026"),
    ("Consolidate Defect Summary", "3", "16/09/2026", "17/09/2026"),
    ("Prepare Final Test Summary", "2", "18/09/2026", "18/09/2026"),
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.6)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(1.7)
section.right_margin = Cm(1.7)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("TEST MILESTONES")
run.bold = True
run.font.name = "Times New Roman"
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0, 0, 0)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
table.autofit = False
widths = [Cm(9.1), Cm(2.7), Cm(3.6), Cm(3.6)]
headers = ["Milestone", "Effort (MD)", "Start Date", "End Date"]

for idx, cell in enumerate(table.rows[0].cells):
    cell.width = widths[idx]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_margins(cell)
    shade(cell, "FCE4D6")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(headers[idx])
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)

for row_index, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for col_index, value in enumerate(values):
        cell = cells[col_index]
        cell.width = widths[col_index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_margins(cell)
        if row_index % 2 == 0:
            shade(cell, "F7F7F7")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(value)
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)

total_cells = table.add_row().cells
total_cells[0].text = "Total"
total_cells[1].text = str(sum(int(row[1]) for row in rows))
total_cells[2].text = ""
total_cells[3].text = ""
for index, cell in enumerate(total_cells):
    cell.width = widths[index]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_margins(cell)
    shade(cell, "D9EAF7")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    for r in p.runs:
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)

doc.save(OUTPUT)
print(OUTPUT)
