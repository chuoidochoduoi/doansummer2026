from copy import copy
from pathlib import Path
import openpyxl

src = Path(r"C:\Users\Administrator\Downloads\CareS_Report_5.1_Unit_Test_Appointment_Review.xlsx")
dst = Path(r"C:\Users\Administrator\Downloads\CareS_Report_5.1_Unit_Test_Appointment_Review_v2.xlsx")
wb = openpyxl.load_workbook(src)

for name, count in (("Appointment 1", 14), ("Appointment 2", 9)):
    ws = wb[name]
    ws.row_dimensions[9].height = 62
    for col in range(6, 6 + count):
        cell = ws.cell(9, col)
        cell.alignment = copy(cell.alignment)
        cell.alignment = openpyxl.styles.Alignment(
            horizontal="center", vertical="center", text_rotation=90,
            wrap_text=False, shrink_to_fit=False
        )
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 5.5

    executed_row = next(
        row for row in range(1, ws.max_row + 1)
        if ws.cell(row, 2).value == "Executed Date"
    )
    for col in range(6, 6 + count):
        ws.cell(executed_row, col).number_format = "dd/mm/yyyy"
        ws.cell(executed_row, col).alignment = openpyxl.styles.Alignment(
            horizontal="center", vertical="center", text_rotation=90
        )
    ws.row_dimensions[executed_row].height = 66
    ws.column_dimensions["B"].width = 48

wb.save(dst)
print(dst)
