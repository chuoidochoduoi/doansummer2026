from copy import copy
from pathlib import Path
import openpyxl

path = Path(r"D:\gitlap\doAnSummer2026\outputs\019fc0c3-dee6-7a11-9b42-26e3805728ed\CareS_Report_5.4_Acceptance_Test_v2.xlsx")
wb = openpyxl.load_workbook(path)
ws = wb["Acceptance Test"]

groups = {
    10: "Global UI, Public Pages & Usability",
    19: "Registration, Authentication & Password Recovery",
    30: "Customer Profile & Family Members",
    39: "Appointments & Check-in",
    51: "Reception, Walk-in Visits & Patient Records",
    62: "Invoices, Payments, Insurance & CareS Card",
    76: "Queue Management & Patient Journey",
    86: "Medical Examination & Prescriptions",
    98: "Laboratory & Paraclinical Workflows",
    109: "Administration, Operations & Support",
}

green = "FF2F751D"
white = "FFFFFFFF"

for row, title in groups.items():
    # Ensure the separator is one visible merged band with no inherited ID formula.
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == row and merged.max_row == row:
            ws.unmerge_cells(str(merged))
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    cell = ws.cell(row, 2)
    cell.value = title
    cell.font = copy(ws["B10"].font)
    cell.fill = copy(ws["B10"].fill)
    cell.border = copy(ws["B10"].border)
    cell.alignment = copy(ws["B10"].alignment)
    cell.font = copy(cell.font)
    cell.font = cell.font.copy(color=white, bold=True)
    cell.fill = copy(cell.fill)
    cell.fill = cell.fill.copy(fill_type="solid", fgColor=green)
    cell.alignment = cell.alignment.copy(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 22

test_rows = [row for row in range(10, 130) if row not in groups]
assert len(test_rows) == 110

for test_id, row in enumerate(test_rows, start=1):
    ws.cell(row, 2).value = test_id
    ws.cell(row, 7).value = True

# Keep statistics limited to actual test rows. Category rows are excluded because G is blank.
ws["B7"] = "=COUNTA(G11:G129)"
ws["G7"] = "=COUNTA(G11:G129)"
ws["H7"] = "=B7-G7"

stats = wb["Test Statistics"]
stats["D11"] = "=COUNTA('Acceptance Test'!G11:G129)"
stats["E11"] = "=F11-D11"
stats["F11"] = "=COUNTA('Acceptance Test'!G11:G129)"
stats["D13"] = "=D11"
stats["E13"] = "=E11"
stats["F13"] = "=F11"
stats["E15"] = "=IFERROR((D13+E13)*100/F13,0)"
stats["E16"] = "=IFERROR(D13*100/F13,0)"

wb.save(path)
print(path)
