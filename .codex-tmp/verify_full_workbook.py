from pathlib import Path
import openpyxl

p = Path(r"C:\Users\Administrator\Downloads\CareS_Report_5.1_Unit_Test_Full_v1.xlsx")
wb = openpyxl.load_workbook(p, data_only=False, read_only=False)
fn = wb["Functions"]
st = wb["Statistics"]
function_sheets = wb.sheetnames[5:]

passed = failed = untested = utc_total = 0
bad_rotation = []
bad_authors = []
bad_sections = []
formula_errors = []
for name in function_sheets:
    ws = wb[name]
    passed += int(ws["B7"].value or 0)
    failed += int(ws["F7"].value or 0)
    untested += int(ws["J7"].value or 0)
    count = int(ws["P7"].value or 0)
    utc_total += count
    if any(ws.cell(9, c).alignment.textRotation != 90 for c in range(6, 6 + count)):
        bad_rotation.append(name)
    if ws["F3"].value != "CuongND" or ws["N3"].value != "CuongND":
        bad_authors.append(name)
    labels = {ws.cell(r, 1).value for r in range(10, ws.max_row + 1)}
    if not {"Condition", "Confirm", "Result"}.issubset(labels):
        bad_sections.append(name)
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "#REF!" in cell.value:
                formula_errors.append(f"{name}!{cell.coordinate}")

for ws in (fn, st):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "#REF!" in cell.value:
                formula_errors.append(f"{ws.title}!{cell.coordinate}")

print({
    "sheets": len(wb.sheetnames),
    "function_sheets": len(function_sheets),
    "utc_total": utc_total,
    "passed": passed,
    "failed": failed,
    "untested": untested,
    "functions_last_id": fn.cell(135, 1).value,
    "statistics_last_id": st.cell(136, 1).value,
    "statistics_total_formula": st.cell(137, 4).value,
    "charts": [(c.anchor._from.row + 1, c.anchor._from.col + 1) for c in st._charts],
    "bad_rotation": bad_rotation[:5],
    "bad_authors": bad_authors[:5],
    "bad_sections": bad_sections[:5],
    "formula_errors": formula_errors[:5],
})
wb.close()
