from pathlib import Path
import json, zipfile, re
import openpyxl

p=Path(r"C:\Users\Administrator\Downloads\CareS_Report_5.1_Unit_Test_Full_v1.xlsx")
wb=openpyxl.load_workbook(p, read_only=True, data_only=False)
names=wb.sheetnames
fs=names[5:]
passed=failed=untested=total=0
bad_rot=[]; bad_author=[]; bad_sections=[]; refs=[]
for n in fs:
    w=wb[n]
    cnt=int(w["P7"].value or 0); total+=cnt
    passed+=int(w["B7"].value or 0); failed+=int(w["F7"].value or 0); untested+=int(w["J7"].value or 0)
    if w["F3"].value!="CuongND" or w["N3"].value!="CuongND": bad_author.append(n)
    if any(w.cell(9,c).alignment.textRotation!=90 for c in range(6,6+cnt)): bad_rot.append(n)
    labels={w.cell(r,1).value for r in range(10,min(w.max_row,120)+1)}
    if not {"Condition","Confirm","Result"}.issubset(labels): bad_sections.append(n)
fn=wb["Functions"]; st=wb["Statistics"]
for w in (fn,st):
    for row in w.iter_rows(min_row=1, max_row=min(w.max_row, 160), min_col=1, max_col=min(w.max_column, 20)):
        for c in row:
            if isinstance(c.value,str) and "#REF!" in c.value: refs.append(f"{w.title}!{c.coordinate}")
result={"sheets":len(names),"function_sheets":len(fs),"utc_total":total,"passed":passed,"failed":failed,"untested":untested,"functions_last_id":fn.cell(135,1).value,"statistics_last_id":st.cell(136,1).value,"statistics_total_formula":st.cell(137,4).value,"bad_rotation":bad_rot,"bad_authors":bad_author,"bad_sections":bad_sections,"formula_errors":refs}
wb.close()
with zipfile.ZipFile(p) as z:
    result["chart_files"]=[n for n in z.namelist() if n.startswith("xl/charts/chart")]
Path('.codex-tmp/full_verify.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')

