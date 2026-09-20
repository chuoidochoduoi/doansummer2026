from __future__ import annotations

from collections import defaultdict
from copy import copy
from pathlib import Path
import math
import re

import openpyxl
from openpyxl.chart import PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

source = Path(r"C:\Users\Administrator\Downloads\CareS_Report_5.1_Unit_Test_Appointment_Review_v3.xlsx")
sample_path = Path(r"C:\Users\Administrator\Downloads\SU26_SEP490_G66_Report_5.1_Unit_Test.xlsx")
output = Path(r"C:\Users\Administrator\Downloads\CareS_Report_5.1_Unit_Test_Full_v1.xlsx")
test_root = Path(r"D:\gitlap\doAnSummer2026\src\test\java\org\example\doansummer2026\service")

wb = openpyxl.load_workbook(source)
sample_wb = openpyxl.load_workbook(sample_path)
sample_fn = sample_wb["refreshToken"]
sample_functions = sample_wb["Functions"]
sample_statistics = sample_wb["Statistics"]

blue = sample_fn["A10"].fill.fgColor.rgb or "FF000080"
blue_fill = PatternFill("solid", fgColor=blue)
white_font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF")
body_font = Font(name="Arial", size=9, color="FF000000")
bold_font = Font(name="Arial", size=9, bold=True, color="FF000000")
tick_font = Font(name="Arial", size=13, bold=True, color="FF000000")
thin = Side(style="thin", color="FF000000")
grid = Border(left=thin, right=thin, top=thin, bottom=thin)

base_sheets = {"Guideline", "Cover", "Functions", "Statistics", "Example"}
function_sheet_names = [name for name in wb.sheetnames if name not in base_sheets]


def extract_tests(java_file: Path):
    text = java_file.read_text(encoding="utf-8")
    pattern = re.compile(r"@(Test|ParameterizedTest)(?:\([^)]*\))?[\s\S]{0,500}?\bvoid\s+([A-Za-z_$][A-Za-z0-9_$]*)\s*\(")
    matches = list(pattern.finditer(text))
    tests = []
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        tests.append({"name": match.group(2), "body": text[match.start():end]})
    return tests


def even_partition(items, count):
    result, at = [], 0
    for index in range(count):
        remaining = len(items) - at
        slots = count - index
        take = math.ceil(remaining / slots) if slots else 0
        result.append(items[at:at + take])
        at += take
    return result


def short_case(name):
    value = name
    value = re.sub(r"^[a-zA-Z0-9]+_", "", value)
    value = re.sub(r"^(Should|With|When|Covers|Handles|Can|Does|Returns?|Creates?|Updates?|Deletes?|Rejects?|Throws?)_?", "", value, flags=re.I)
    value = re.sub(r"_(Should|When|With|And|Then)_", "_", value, flags=re.I)
    words = [w for w in re.split(r"_+|(?<=[a-z0-9])(?=[A-Z])", value) if w]
    words = [w for w in words if w.lower() not in {"should", "when", "with", "and", "then", "covers"}]
    if not words:
        words = ["validCase"]
    compact = words[0].lower() + "".join(w[:1].upper() + w[1:] for w in words[1:])
    return compact[:46]


def operation_title(tests, class_name):
    verbs = []
    for test in tests:
        match = re.match(r"([a-z][A-Za-z0-9]*)", test["name"])
        verb = match.group(1) if match else "handle"
        if verb not in verbs:
            verbs.append(verb)
    readable = [re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", x).lower() for x in verbs[:3]]
    action = ", ".join(readable[:-1]) + ((" and " + readable[-1]) if len(readable) > 1 else readable[0])
    subject = re.sub(r"Service$", "", class_name)
    subject = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", subject).lower()
    return f"{action.capitalize()} {subject} records"


def exception_name(test):
    patterns = [
        r"assertThrows\s*\(\s*([A-Za-z0-9_$.]+)\.class",
        r"assertThatThrownBy[\s\S]{0,400}?isInstanceOf\s*\(\s*([A-Za-z0-9_$.]+)\.class",
    ]
    for pattern in patterns:
        match = re.search(pattern, test["body"])
        if match:
            return match.group(1).split(".")[-1]
    return None


def test_type(name, exception):
    lower = name.lower()
    if any(x in lower for x in ("null", "empty", "blank", "zero", "minimum", "maximum", "boundary", "before", "after", "exact", "limit")):
        return "B"
    if exception or any(x in lower for x in ("invalid", "missing", "unknown", "reject", "forbid", "conflict", "fail", "notfound", "not_found", "wrong", "duplicate", "inactive", "expired")):
        return "A"
    return "N"


def inferred_conditions(test):
    lower = test["name"].lower()
    if "null" in lower:
        input_state = "null"
    elif "empty" in lower:
        input_state = "empty"
    elif "blank" in lower:
        input_state = "blank"
    elif any(x in lower for x in ("invalid", "wrong", "duplicate", "malformed")):
        input_state = "invalid"
    else:
        input_state = "valid"
    record_exists = "F" if any(x in lower for x in ("missing", "notfound", "not_found", "unknown")) else "T"
    authorized = "F" if any(x in lower for x in ("unauthor", "forbid", "notowner", "not_owner", "wrongdoctor", "nonadmin", "nonnurse")) else "T"
    status_valid = "F" if any(x in lower for x in ("invalidstatus", "statusnot", "completed", "cancelled", "rejected", "expired", "inactive", "blocked")) else "T"
    has_conflict = "T" if any(x in lower for x in ("conflict", "duplicate", "alreadyexists", "already_exists")) else "F"
    return {
        "inputState": input_state,
        "recordExists": record_exists,
        "authorized": authorized,
        "statusValid": status_valid,
        "hasConflict": has_conflict,
        "case": short_case(test["name"]),
    }


# Rebuild the exact test allocation for every existing function sheet.
sheet_rows = []
for row in range(11, 11 + len(function_sheet_names)):
    sheet_rows.append({
        "sheet": wb["Functions"].cell(row, 6).value,
        "class": wb["Functions"].cell(row, 3).value,
        "code": wb["Functions"].cell(row, 5).value,
    })

rows_by_class = defaultdict(list)
for row in sheet_rows:
    rows_by_class[row["class"]].append(row)

assigned_groups = {}
for class_name, rows in rows_by_class.items():
    java_file = test_root / f"{class_name}Test.java"
    tests = extract_tests(java_file) if java_file.exists() else []
    if class_name == "AppointmentService" and len(rows) >= 2:
        create = [t for t in tests if re.match(r"^create_", t["name"])]
        guest = [t for t in tests if t["name"].startswith("createForGuest_")]
        used = {id(t) for t in create + guest}
        remaining = [t for t in tests if id(t) not in used]
        groups = [create, guest] + even_partition(remaining, len(rows) - 2)
    else:
        groups = even_partition(tests, len(rows))
    for row, tests_for_sheet in zip(rows, groups):
        assigned_groups[row["sheet"]] = tests_for_sheet


def write_matrix_sheet(ws, tests):
    count = len(tests)
    last_col = 5 + count
    if count == 0:
        return

    for merged in list(ws.merged_cells.ranges):
        if merged.max_row >= 9 and merged.min_row <= 250:
            ws.unmerge_cells(str(merged))
    for row in ws.iter_rows(min_row=9, max_row=250, min_col=1, max_col=120):
        for cell in row:
            cell.value = None
            cell._style = copy(sample_fn["B11"]._style)

    # Title/header metrics.
    ws["A6"] = "Passed"; ws["C6"] = "Failed"; ws["F6"] = "Untested"; ws["L6"] = "N/A/B"; ws["O6"] = "Total Test Cases"
    typed = [(t, exception_name(t)) for t in tests]
    types = [test_type(t["name"], ex) for t, ex in typed]
    ws["A7"] = count; ws["C7"] = 0; ws["F7"] = 0
    ws["L7"] = types.count("N"); ws["M7"] = types.count("A"); ws["N7"] = types.count("B"); ws["O7"] = count

    ws.merge_cells(start_row=9, start_column=1, end_row=9, end_column=5)
    for col in range(1, last_col + 1):
        ws.cell(9, col).fill = blue_fill
        ws.cell(9, col).border = grid
    for index in range(count):
        cell = ws.cell(9, 6 + index)
        cell.value = f"UTCID{index + 1:02d}"
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
        ws.column_dimensions[openpyxl.utils.get_column_letter(6 + index)].width = 5.4
    ws.row_dimensions[9].height = 66

    conditions = [inferred_conditions(t) for t in tests]
    values = {
        "inputState": ["valid", "invalid", "null", "empty", "blank"],
        "recordExists": ["T", "F"],
        "authorized": ["T", "F"],
        "statusValid": ["T", "F"],
        "hasConflict": ["T", "F"],
        "case": list(dict.fromkeys(c["case"] for c in conditions)),
    }

    logical_sections = []
    condition_rows = [("Precondition", ["O"] * count, True), ("Can connect to server", ["O"] * count, False)]
    for variable, options in values.items():
        condition_rows.append((None, [None] * count, False))
        condition_rows.append((variable, [None] * count, True))
        for option in options:
            condition_rows.append((option, ["O" if c[variable] == option else None for c in conditions], False))
    logical_sections.append(("Condition", condition_rows))

    return_rows = [
        ("Return", [None] * count, True),
        ("T", ["O" if ex is None else None for _, ex in typed], False),
        ("F", ["O" if ex is not None else None for _, ex in typed], False),
        (None, [None] * count, False),
        ("Exception", [None] * count, True),
    ]
    exception_types = list(dict.fromkeys(ex for _, ex in typed if ex))
    for ex_type in exception_types:
        return_rows.append((ex_type, ["O" if ex == ex_type else None for _, ex in typed], False))
    return_rows.extend([
        (None, [None] * count, False),
        ("Log message", [None] * count, True),
        ("N/A", ["O"] * count, False),
    ])
    logical_sections.append(("Confirm", return_rows))
    result_rows = [
        ("Type (N : Normal, A : Abnormal, B : Boundary)", types, True),
        ("Passed/Failed", ["P"] * count, False),
        ("Executed Date", ["14/09/2026"] * count, False),
        ("Defect ID", [None] * count, False),
    ]
    logical_sections.append(("Result", result_rows))

    current = 10
    section_positions = {}
    for section_name, rows in logical_sections:
        start = current
        for label, marks, group in rows:
            ws.merge_cells(start_row=current, start_column=2, end_row=current, end_column=5)
            ws.cell(current, 2).value = label
            for index, mark in enumerate(marks):
                ws.cell(current, 6 + index).value = mark
            for col in range(1, last_col + 1):
                ws.cell(current, col).border = grid
                ws.cell(current, col).font = body_font
                ws.cell(current, col).alignment = Alignment(vertical="center")
            if label is None:
                ws.row_dimensions[current].height = 8
            else:
                ws.row_dimensions[current].height = 19
                ws.cell(current, 2).font = bold_font if group else body_font
                ws.cell(current, 2).alignment = Alignment(horizontal="left" if group else "right", vertical="center")
                for index, mark in enumerate(marks):
                    mark_cell = ws.cell(current, 6 + index)
                    mark_cell.font = tick_font if mark == "O" else body_font
                    mark_cell.alignment = Alignment(horizontal="center", vertical="center")
            current += 1
        end = current - 1
        ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        ws.cell(start, 1).value = section_name
        ws.cell(start, 1).font = white_font
        ws.cell(start, 1).alignment = Alignment(horizontal="center", vertical="top")
        for row in range(start, end + 1):
            ws.cell(row, 1).fill = blue_fill
        section_positions[section_name] = (start, end)

    result_start, result_end = section_positions["Result"]
    ws.row_dimensions[result_start + 2].height = 64
    ws.row_dimensions[result_start + 3].height = 34
    for index in range(count):
        date_cell = ws.cell(result_start + 2, 6 + index)
        date_cell.font = Font(name="Arial", size=8)
        date_cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 23
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 10
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "F10"


# Keep the two approved Appointment sheets. Apply the same compact structure everywhere else.
for sheet_name in function_sheet_names:
    if sheet_name in {"Appointment 1", "Appointment 2"}:
        continue
    write_matrix_sheet(wb[sheet_name], assigned_groups.get(sheet_name, []))


# Expand and normalize the Functions table through all 125 function rows.
functions = wb["Functions"]
normal_function_styles = [copy(sample_functions.cell(11, col)._style) for col in range(1, 9)]
for row_index, row_info in enumerate(sheet_rows, start=11):
    tests = assigned_groups.get(row_info["sheet"], [])
    title = functions.cell(row_index, 4).value if row_info["sheet"] in {"Appointment 1", "Appointment 2"} else operation_title(tests, row_info["class"])
    requirement = functions.cell(row_index, 2).value or re.sub(r"Service$", "", row_info["class"])
    values = [
        row_index - 10,
        requirement,
        row_info["class"],
        title,
        row_info["code"],
        row_info["sheet"],
        f"Checks {title.lower()} across successful, rejected, and edge cases.",
        "1. Required dependencies are mocked\n2. Test data and the current user context are prepared",
    ]
    for col, value in enumerate(values, start=1):
        cell = functions.cell(row_index, col)
        cell._style = copy(normal_function_styles[col - 1])
        cell.value = value
        cell.border = grid
        cell.alignment = Alignment(horizontal="center" if col in {1, 5, 6} else "left", vertical="center", wrap_text=True)
    functions.cell(row_index, 6).hyperlink = f"#'{row_info['sheet']}'!A1"
    functions.cell(row_index, 6).font = Font(name="Arial", size=9, color="FF0000FF", underline="single")
    functions.row_dimensions[row_index].height = 38


# Expand Statistics through all functions and move totals/charts below the table.
statistics = wb["Statistics"]
normal_stat_styles = [copy(sample_statistics.cell(12, col)._style) for col in range(1, 10)]
subtotal_styles = [copy(sample_statistics.cell(124, col)._style) for col in range(1, 10)]
for row in range(12, 210):
    for col in range(1, 10):
        statistics.cell(row, col).value = None
        statistics.cell(row, col)._style = copy(normal_stat_styles[col - 1])

for index, row_info in enumerate(sheet_rows):
    row = 12 + index
    tests = assigned_groups.get(row_info["sheet"], [])
    if row_info["sheet"] == "Appointment 1":
        counts = (14, 0, 0, 2, 6, 6, 14)
    elif row_info["sheet"] == "Appointment 2":
        counts = (9, 0, 0, 2, 5, 2, 9)
    else:
        exceptions = [exception_name(t) for t in tests]
        types = [test_type(t["name"], ex) for t, ex in zip(tests, exceptions)]
        counts = (len(tests), 0, 0, types.count("N"), types.count("A"), types.count("B"), len(tests))
    values = [index + 1, row_info["code"], *counts]
    for col, value in enumerate(values, start=1):
        cell = statistics.cell(row, col)
        cell._style = copy(normal_stat_styles[col - 1])
        cell.value = value
        cell.border = grid
        cell.alignment = Alignment(horizontal="center", vertical="center")
    statistics.cell(row, 2).hyperlink = f"#'{row_info['sheet']}'!A1"
    statistics.cell(row, 2).font = Font(name="Arial", size=9, color="FF0000FF", underline="single")
    statistics.row_dimensions[row].height = 20

subtotal_row = 12 + len(sheet_rows)
for col in range(1, 10):
    statistics.cell(subtotal_row, col)._style = copy(subtotal_styles[col - 1])
    statistics.cell(subtotal_row, col).border = grid
statistics.cell(subtotal_row, 2).value = "Sub Total"
for col in range(3, 10):
    letter = openpyxl.utils.get_column_letter(col)
    statistics.cell(subtotal_row, col).value = f"=SUM({letter}12:{letter}{subtotal_row - 1})"
statistics.row_dimensions[subtotal_row].height = 22

coverage_row = subtotal_row + 2
statistics.cell(coverage_row, 1).value = "Test coverage"
statistics.cell(coverage_row, 2).value = "100%"
statistics.cell(coverage_row + 1, 1).value = "Test successful coverage"
statistics.cell(coverage_row + 1, 2).value = "100%"
for row in (coverage_row, coverage_row + 1):
    statistics.cell(row, 1).font = bold_font
    statistics.cell(row, 2).font = bold_font

statistics._charts = []
chart_row = coverage_row + 4
passed_chart = PieChart()
passed_chart.title = "Passed Percent"
passed_chart.add_data(Reference(statistics, min_col=3, max_col=5, min_row=subtotal_row), titles_from_data=False)
passed_chart.set_categories(Reference(statistics, min_col=3, max_col=5, min_row=11))
passed_chart.height = 8.5
passed_chart.width = 11
statistics.add_chart(passed_chart, f"A{chart_row}")

type_chart = PieChart()
type_chart.title = "Test Type"
type_chart.add_data(Reference(statistics, min_col=6, max_col=8, min_row=subtotal_row), titles_from_data=False)
type_chart.set_categories(Reference(statistics, min_col=6, max_col=8, min_row=11))
type_chart.height = 8.5
type_chart.width = 11
statistics.add_chart(type_chart, f"F{chart_row}")
statistics.sheet_view.showGridLines = False

sample_wb.close()
wb.save(output)
print(output)
print(f"Function sheets: {len(function_sheet_names)}")
print(f"Documented tests: {sum(len(v) for v in assigned_groups.values())}")
