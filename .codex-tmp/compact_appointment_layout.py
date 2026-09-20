from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

src = Path(r"C:\Users\Administrator\Downloads\CareS_Report_5.1_Unit_Test_Appointment_Review_v2.xlsx")
sample_path = Path(r"C:\Users\Administrator\Downloads\SU26_SEP490_G66_Report_5.1_Unit_Test.xlsx")
dst = Path(r"C:\Users\Administrator\Downloads\CareS_Report_5.1_Unit_Test_Appointment_Review_v3.xlsx")

wb = openpyxl.load_workbook(src)
sample_wb = openpyxl.load_workbook(sample_path)
sample = sample_wb["refreshToken"]
blue = sample["A10"].fill.fgColor.rgb or "FF000080"
blue_fill = PatternFill("solid", fgColor=blue)
white_font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF")
body_font = Font(name="Arial", size=9, color="FF000000")
bold_font = Font(name="Arial", size=9, bold=True, color="FF000000")
tick_font = Font(name="Arial", size=13, bold=True, color="FF000000")
thin = Side(style="thin", color="FF000000")
grid = Border(left=thin, right=thin, top=thin, bottom=thin)

condition_groups = {
    "Appointment 1": {"Precondition", "accountExists", "role", "profileExists", "serviceIds", "hasConflict",
                       "shiftExists", "serviceExists", "dateOfBirth", "ageEligible", "gender",
                       "allowCustomerBooking", "scheduledAt"},
    "Appointment 2": {"Precondition", "guestFullName", "guestPhone", "guestEmail", "guestAge", "guestGender",
                       "scheduledAt", "shiftId", "serviceIds", "allowCustomerBooking", "hasGuestConflict",
                       "conflictMatchedBy"},
}
confirm_groups = {"Return", "status", "customerLinked", "servicesLinked", "isGuest", "customerId",
                  "guestInformationSaved", "repository.save", "notifyReceptionists", "Exception", "Log message"}

def col_letter(index):
    return openpyxl.utils.get_column_letter(index)

def read_logical_rows(ws, start, end, count):
    rows = []
    for r in range(start, end + 1):
        rows.append({
            "a": ws.cell(r, 1).value,
            "b": ws.cell(r, 2).value,
            "marks": [ws.cell(r, 6 + i).value for i in range(count)],
        })
    return rows

def compact(rows, groups):
    out = []
    for item in rows:
        label = item["b"]
        if label in groups and out and out[-1]["b"] is not None:
            out.append({"a": None, "b": None, "marks": [None] * len(item["marks"]), "spacer": True})
        item = dict(item)
        item["group"] = label in groups
        out.append(item)
    return out

for name in ("Appointment 1", "Appointment 2"):
    ws = wb[name]
    count = int(ws["O7"].value)
    last_col = 5 + count
    old_result = next(r for r in range(10, 180) if ws.cell(r, 1).value == "Result")
    old_confirm = next(r for r in range(10, old_result) if ws.cell(r, 1).value == "Confirm")
    condition = compact(read_logical_rows(ws, 10, old_confirm - 1, count), condition_groups[name])
    confirm = compact(read_logical_rows(ws, old_confirm, old_result - 1, count), confirm_groups)
    result = read_logical_rows(ws, old_result, old_result + 3, count)

    # Remove all old matrix merges, then clear the old matrix area.
    for merged in list(ws.merged_cells.ranges):
        if merged.max_row >= 9 and merged.min_row <= 180:
            ws.unmerge_cells(str(merged))
    for row in ws.iter_rows(min_row=9, max_row=180, min_col=1, max_col=80):
        for cell in row:
            cell.value = None
            cell._style = copy(sample["B11"]._style)

    # Header: one clean blue band and vertical UTC labels.
    ws.merge_cells(start_row=9, start_column=1, end_row=9, end_column=5)
    for c in range(1, last_col + 1):
        cell = ws.cell(9, c)
        cell.fill = blue_fill
        cell.border = grid
    for i in range(count):
        cell = ws.cell(9, 6 + i)
        cell.value = f"UTCID{i + 1:02d}"
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
        ws.column_dimensions[col_letter(6 + i)].width = 5.4
    ws.row_dimensions[9].height = 66

    current = 10
    section_ranges = []
    for section_name, rows in (("Condition", condition), ("Confirm", confirm), ("Result", result)):
        section_start = current
        for item in rows:
            spacer = item.get("spacer", False)
            ws.merge_cells(start_row=current, start_column=2, end_row=current, end_column=5)
            ws.cell(current, 2).value = item["b"]
            for i, mark in enumerate(item["marks"]):
                ws.cell(current, 6 + i).value = mark
            for c in range(1, last_col + 1):
                ws.cell(current, c).border = grid
                ws.cell(current, c).font = body_font
                ws.cell(current, c).alignment = Alignment(vertical="center")
            if spacer:
                ws.row_dimensions[current].height = 8
            else:
                ws.row_dimensions[current].height = 19
                if item.get("group", False) or current == section_start:
                    ws.cell(current, 2).font = bold_font
                    ws.cell(current, 2).alignment = Alignment(horizontal="left", vertical="center")
                else:
                    ws.cell(current, 2).alignment = Alignment(horizontal="right", vertical="center")
                for i, mark in enumerate(item["marks"]):
                    mark_cell = ws.cell(current, 6 + i)
                    mark_cell.alignment = Alignment(horizontal="center", vertical="center")
                    mark_cell.font = tick_font if mark == "O" else body_font
            current += 1
        section_end = current - 1
        ws.merge_cells(start_row=section_start, start_column=1, end_row=section_end, end_column=1)
        label = ws.cell(section_start, 1)
        label.value = section_name
        label.fill = blue_fill
        label.font = white_font
        label.alignment = Alignment(horizontal="center", vertical="top")
        for r in range(section_start, section_end + 1):
            ws.cell(r, 1).fill = blue_fill
        section_ranges.append((section_name, section_start, section_end))

    # Result rows use compact, consistent text; dates stay readable in narrow UTC columns.
    result_start = next(start for sec, start, end in section_ranges if sec == "Result")
    for i in range(count):
        ws.cell(result_start, 6 + i).font = body_font
        ws.cell(result_start + 1, 6 + i).font = body_font
        ws.cell(result_start + 2, 6 + i).font = Font(name="Arial", size=8)
        ws.cell(result_start + 2, 6 + i).alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
    ws.row_dimensions[result_start + 2].height = 64
    ws.row_dimensions[result_start + 3].height = 34

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 23
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 10
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "F10"

sample_wb.close()
wb.save(dst)
print(dst)
